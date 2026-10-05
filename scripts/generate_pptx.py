import os
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Cores da identidade visual Doce Gestor
    PRIMARY = RGBColor(179, 57, 81)       # Framboesa / Rosé Nobre (#B33951)
    SECONDARY = RGBColor(140, 83, 62)     # Chocolate Suave (#8C533E)
    DARK = RGBColor(43, 37, 35)           # Café Escuro (#2B2523)
    LIGHT_BG = RGBColor(250, 247, 245)    # Creme Suave (#FAF7F5)
    CARD_BG = RGBColor(255, 255, 255)     # Branco Puro
    ACCENT_GOLD = RGBColor(212, 149, 74)  # Caramelo Dourado (#D4954A)
    MUTED = RGBColor(115, 105, 100)       # Cinza Acolhedor (#736964)
    SUCCESS = RGBColor(46, 125, 50)       # Verde Pistache / Sucesso
    BORDER_COLOR = RGBColor(230, 220, 214)

    blank_layout = prs.slide_layouts[6]

    def add_bg(slide, color=LIGHT_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="DOCE GESTOR • SISTEMA DE GESTÃO PARA CONFEITARIAS"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_GOLD

        # Header title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK

    # ==========================================
    # SLIDE 1: CAPA
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, PRIMARY)

    # Accent decorative box
    acc = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.1))
    acc.fill.solid()
    acc.fill.fore_color.rgb = RGBColor(255, 255, 255)
    acc.line.color.rgb = ACCENT_GOLD
    acc.line.width = Pt(2)

    tb = s1.shapes.add_textbox(Inches(1.4), Inches(1.7), Inches(10.5), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "PLANO ESTRATÉGICO & CRONOGRAMA DE IMPLANTAÇÃO"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_GOLD

    p1 = tf.add_paragraph()
    p1.text = "Doce Gestor — Sistema Web Multiempresa"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = PRIMARY
    p1.space_before = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "Especializado para Confeitarias Caseiras e Negócios de Alimentação"
    p2.font.size = Pt(18)
    p2.font.color.rgb = SECONDARY
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "• Arquitetura Multi-Tenant isolada por E-mail e Licença\n• Wireframes das Telas Essenciais (Fichas Técnicas, Cardápio Digital e Pedidos)\n• Opinião Técnica e Passo a Passo para Apresentação às Clientes Piloto"
    p3.font.size = Pt(14)
    p3.font.color.rgb = DARK
    p3.space_before = Pt(20)

    p4 = tf.add_paragraph()
    p4.text = "Visão de Engenharia & Arquitetura de Software | Outubro de 2026"
    p4.font.size = Pt(11)
    p4.font.color.rgb = MUTED
    p4.space_before = Pt(30)

    # ==========================================
    # SLIDE 2: OPINIÃO HONESTA DO ARQUITETO
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Opinião Técnica Honesta: Onde Começar e Por Quê?")

    cards = [
        ("1. O Choque de Realidade do Nicho", 
         "A confeiteira caseira trabalha sozinha ou com um ajudante. Ela não é operadora de ERP. Seu tempo é dividido entre as panelas, o forno e responder clientes no WhatsApp.\n\nSistemas genéricos falham porque exigem burocracia contábil pesada antes de entregar qualquer valor prático.",
         PRIMARY),
        ("2. A Dor Primária: 'Tenho Lucro ou Prejuízo?'", 
         "A dor nº 1 de toda confeiteira é a PRECIFICAÇÃO. Elas compram leite condensado, gastam gás, usam fitas e forminhas, mas vendem o bolo no 'chute' ou olhando o preço da vizinha.\n\nSe o sistema calcular o custo exato por grama e sugerir o preço de venda justo no primeiro dia, o sistema se vende sozinho.",
         ACCENT_GOLD),
        ("3. Diagnóstico do Código Atual", 
         "O projeto atual tem um módulo de Funcionários muito bem desenhado (comissões, histórico, folha). Isso é fantástico para a Fase 2!\n\nPorém, para as CLIENTES PILOTO caseiras, devemos priorizar: Multi-tenant -> Insumos/Ficha Técnica -> Cardápio com Link de Pedido. Funcionários vêm em seguida.",
         SECONDARY),
    ]

    for i, (title, desc, bar_col) in enumerate(cards):
        left = Inches(0.8 + i * 4.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Top color accent
        top_bar = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = bar_col
        top_bar.line.fill.background()

        tb = s2.shapes.add_textbox(left + Inches(0.25), Inches(2.3), Inches(3.2), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = bar_col

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = DARK
        p_d.space_before = Pt(12)

    # ==========================================
    # SLIDE 3: ARQUITETURA MULTIEMPRESA & CONTROLE
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Arquitetura Multi-Tenant: Isolamento e Controle de Licença")

    arch_items = [
        ("Isolamento Rigoroso por E-mail & Tenant ID",
         "• Cada confeitaria possui seu próprio registro em tb_empresa.\n• O e-mail de login é o identificador mestre do ateliê.\n• TODAS as queries no banco filtram obrigatoriamente WHERE empresa_id = ?.\n• Zero risco de vazamento de dados de receitas ou clientes entre confeitarias concorrentes."),
        ("Controle por Chave Serial & Licença",
         "• tb_empresa conterá campos de licença: chave_serial, tipo_plano, status_licenca e data_expiracao.\n• Modelo Inicial: 15 dias de teste grátis (Trial).\n• Ativação definitiva via Serial Gerado pelo Administrador ou renovação por assinatura mensal.\n• Middleware de segurança que bloqueia o acesso caso a licença expire."),
        ("Transição Fluida: SQLite Local ➔ Nuvem Turso",
         "• O backend já utiliza @libsql/client. Isso é uma escolha de engenharia genial!\n• Em desenvolvimento: roda no arquivo local doce-gestor.db sem custos de servidor.\n• Em produção comercial: a mesma biblioteca conecta diretamente ao cluster Turso (LibSQL Serverless) apenas trocando uma variável de ambiente."),
    ]

    for i, (title, desc) in enumerate(arch_items):
        top = Inches(2.0 + i * 1.6)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.733), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        # Left bar
        lbar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(0.15), Inches(1.4))
        lbar.fill.solid()
        lbar.fill.fore_color.rgb = PRIMARY if i == 0 else (ACCENT_GOLD if i == 1 else SECONDARY)
        lbar.line.fill.background()

        tb = s3.shapes.add_textbox(Inches(1.2), top + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = DARK

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11.5)
        pd.font.color.rgb = MUTED
        pd.space_before = Pt(4)

    # ==========================================
    # SLIDE 4: TELA 1 - FICHA TÉCNICA E PRECIFICAÇÃO
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Tela Essencial 1: Insumos, Fichas Técnicas & Formação de Preço")

    # Coluna Esquerda: Mockup conceitual da tela
    mockup_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(7.2), Inches(5.0))
    mockup_box.fill.solid()
    mockup_box.fill.fore_color.rgb = CARD_BG
    mockup_box.line.color.rgb = BORDER_COLOR

    tb_m = s4.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(6.8), Inches(4.7))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p = tf_m.paragraphs[0]
    p.text = "🎂 RECEITA: Bolo Red Velvet com Ninho (Rendimento: 1 Bolo 1.8kg)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p = tf_m.add_paragraph()
    p.text = "───────────────────────────────────────────────────────\n" \
             "Ingrediente            Qtd Usada     Embalagem/Preço      Custo Real\n" \
             "Leite Condensado       395g          395g  - R$ 6,50      R$ 6,50\n" \
             "Farinha de Trigo       300g          1000g - R$ 5,00      R$ 1,50\n" \
             "Cream Cheese           250g          300g  - R$ 14,00     R$ 11,66\n" \
             "Manteiga Extra         100g          200g  - R$ 12,00     R$ 6,00\n" \
             "Embalagem / Cakeboard  1 un          1 un  - R$ 4,50      R$ 4,50\n" \
             "───────────────────────────────────────────────────────\n" \
             "Subtotal Ingredientes + Embalagem:                      R$ 30,16\n" \
             "Custos Fixos (Gás, Energia, Água - 15%):                R$  4,52\n" \
             "Mão de Obra da Confeiteira (1.5h a R$ 20/h):            R$ 30,00\n" \
             "═══════════════════════════════════════════════════════\n" \
             "CUSTO TOTAL DE PRODUÇÃO:                                R$ 64,68\n" \
             "Margem de Lucro Desejada (100% s/ custo):              +R$ 64,68\n" \
             "PREÇO SUGERIDO DE VENDA:                                R$ 129,36"
    p.font.size = Pt(10)
    p.font.name = "Consolas"
    p.font.color.rgb = DARK

    # Coluna Direita: Benefícios para as Piloto
    right_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = CARD_BG
    right_box.line.color.rgb = BORDER_COLOR

    tb_r = s4.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "O que essa tela entrega para a confeiteira?"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    p = tf_r.add_paragraph()
    p.text = "\n1. Fim das perdas invisíveis:\nO sistema converte gramas, quilos, litros e fardos automaticamente.\n\n" \
             "2. Atualização em cascata:\nSe o leite condensado subir no mercado, a confeiteira atualiza o insumo e TODOS os bolos recalculam o custo sozinhos!\n\n" \
             "3. Valorização do trabalho:\nInclui o valor da hora dela na conta. Ela passa a receber salário além do lucro do negócio."
    p.font.size = Pt(11.5)
    p.font.color.rgb = DARK

    # ==========================================
    # SLIDE 5: TELA 2 - CARDÁPIO DIGITAL PÚBLICO
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Tela Essencial 2: Cardápio Digital Público & Pedido via Link")

    # Coluna Esquerda: Mockup Mobile do Cardápio
    mock_mob = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(4.6), Inches(5.0))
    mock_mob.fill.solid()
    mock_mob.fill.fore_color.rgb = CARD_BG
    mock_mob.line.color.rgb = BORDER_COLOR

    tb_mob = s5.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(4.2), Inches(4.7))
    tf_mob = tb_mob.text_frame
    tf_mob.word_wrap = True

    p = tf_mob.paragraphs[0]
    p.text = "📱 docegestor.com.br/atelie-doce-afeto\n"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = MUTED

    p = tf_mob.add_paragraph()
    p.text = "🧁 Ateliê Açúcar & Afeto\nBolos artesanais, doces e kits festa com amor.\n"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p = tf_mob.add_paragraph()
    p.text = "Categorias: [Bolos] [Doces Finos] [Kits Festa]\n" \
             "───────────────────────────────────\n" \
             "🎂 Bolo Vulcão Ninho c/ Nutella   R$ 65,00\n" \
             "   Massa fofinha de chocolate e muito recheio.\n" \
             "   [ + Adicionar ao Pedido ]\n\n" \
             "🍓 Copo da Felicidade Morango     R$ 22,00\n" \
             "   Brigadeiro branco, geleia e brownie.\n" \
             "   [ + Adicionar ao Pedido ]\n" \
             "───────────────────────────────────\n" \
             "🛒 Seu Carrinho: 2 itens (R$ 87,00)\n" \
             "📅 Data p/ Retirada: Sábado, 15:00\n" \
             "👉 [ ENVIAR PEDIDO NO WHATSAPP ]"
    p.font.size = Pt(9.5)
    p.font.name = "Consolas"
    p.font.color.rgb = DARK

    # Coluna Direita: O Fluxo Automático
    flow_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.7), Inches(1.9), Inches(6.8), Inches(5.0))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = CARD_BG
    flow_box.line.color.rgb = BORDER_COLOR

    tb_fl = s5.shapes.add_textbox(Inches(6.0), Inches(2.1), Inches(6.2), Inches(4.6))
    tf_fl = tb_fl.text_frame
    tf_fl.word_wrap = True

    p = tf_fl.paragraphs[0]
    p.text = "Como funciona o fluxo de pedidos das clientes piloto?"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GOLD

    steps = [
        ("1. Divulgação Simples:", "A confeiteira coloca o link na bio do Instagram e no status do WhatsApp."),
        ("2. Autoatendimento sem Atrito:", "O cliente entra pelo celular, vê fotos reais, preços, escolhe recheios e não precisa instalar nenhum app."),
        ("3. Mensagem Pronta no WhatsApp:", "Ao finalizar, o sistema abre o WhatsApp da confeiteira com a mensagem 100% formatada: itens, adicionais, data, horário e valor total."),
        ("4. Registro Automático no Painel:", "O pedido cai automaticamente na tela de encomendas do sistema com status 'Pendente / Orçamento', pronto para a confeiteira confirmar.")
    ]
    for s_title, s_desc in steps:
        p = tf_fl.add_paragraph()
        p.text = f"{s_title} {s_desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 6: TELA 3 - GESTÃO DE PEDIDOS & AGENDA
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Tela Essencial 3: Quadro de Encomendas & Agenda de Produção")

    kanban_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.733), Inches(5.0))
    kanban_box.fill.solid()
    kanban_box.fill.fore_color.rgb = CARD_BG
    kanban_box.line.color.rgb = BORDER_COLOR

    tb_k = s6.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.3), Inches(4.6))
    tf_k = tb_k.text_frame
    tf_k.word_wrap = True

    p = tf_k.paragraphs[0]
    p.text = "📋 Kanban de Produção & Entregas da Semana"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    cols_text = (
        "┌────────────────────┬────────────────────┬────────────────────┬────────────────────┐\n"
        "│ 📝 NOVO ORÇAMENTO  │ 🧁 CONFIRMADO (50%)│ 🎂 EM PRODUÇÃO     │ ✨ PRONTO / ENTREGUE│\n"
        "├────────────────────┼────────────────────┼────────────────────┼────────────────────┤\n"
        "│ #104 - Mariana S.  │ #102 - Camila Lima │ #101 - Festa Lucas │ #099 - Carlos M.   │\n"
        "│ Bolo 2kg Red Velvet│ 100 Brigadeiros    │ Kit Festa 30 pess. │ Copo Felicidade    │\n"
        "│ Entrega: Sáb 14h   │ Entrega: Sex 18h   │ Entrega: Sáb 10h   │ Entregue hoje      │\n"
        "│ Valor: R$ 130,00   │ Sinal R$ 75 PAGO   │ Sinal R$ 150 PAGO  │ R$ 22,00 RECEBIDO  │\n"
        "│ [Confirmar Sinal]  │ [Iniciar Produção] │ [Avisar no Whats]  │ [Arquivar]         │\n"
        "└────────────────────┴────────────────────┴────────────────────┴────────────────────┘\n\n"
        "⭐ Recursos Críticos para Confeiteiras:\n"
        "• Alerta de Sinal/Entrada: Destaca visualmente se o cliente pagou os 50% para reservar a data.\n"
        "• Lista de Compras da Semana: O sistema soma os ingredientes de todos os bolos confirmados e gera a lista do supermercado!\n"
        "• Botão 1-Clique WhatsApp: 'Seu bolo está pronto para retirada!'"
    )
    p = tf_k.add_paragraph()
    p.text = cols_text
    p.font.size = Pt(10)
    p.font.name = "Consolas"
    p.font.color.rgb = DARK
    p.space_before = Pt(8)

    # ==========================================
    # SLIDE 7: CRONOGRAMA DE EXECUÇÃO
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Cronograma de Execução: Da Base ao Lançamento Piloto")

    timeline = [
        ("FASE 0 • Semana 1", "Fundação Multi-Tenant & Acesso", 
         "• Cadastro self-service da Confeitaria (Nome, WhatsApp, E-mail).\n• Login seguro por e-mail e isolamento estrito de dados.\n• Controle de Licença (Trial 15 dias / Chave Serial)."),
        ("FASE 1 • Semana 2", "Insumos, Receitas & Precificação", 
         "• Cadastro de insumos com conversão de unidades (g, kg, ml).\n• Ficha técnica com cálculo de custo por ingrediente + gás/embalagem.\n• Simulador de margem de lucro e preço de venda sugerido."),
        ("FASE 2 • Semana 3", "Cardápio Digital Público & WhatsApp", 
         "• Página pública do cardápio responsiva para celular (/cardapio/nome).\n• Carrinho de encomendas sem necessidade de login do cliente.\n• Integração direta para fechar pedido no WhatsApp da confeitaria."),
        ("FASE 3 • Semana 4", "Gestão de Encomendas & Produção", 
         "• Quadro Kanban de pedidos (Orçamento -> Produção -> Pronto).\n• Agenda de entregas por dia/hora com controle de sinal (50%).\n• Lista automática de compras de ingredientes para a semana."),
        ("FASE 4 • Semana 5", "Financeiro & Equipe (Conexão do Atual)", 
         "• Fluxo de caixa simples (Entradas e Saídas do ateliê).\n• Integração com o módulo já pronto de Funcionários, comissões e folha.\n• Painel de lucros mensais consolidados."),
    ]

    for i, (fase, title, desc) in enumerate(timeline):
        top = Inches(1.9 + i * 1.05)
        bar = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.733), Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = CARD_BG
        bar.line.color.rgb = BORDER_COLOR

        # Fase badge
        badge = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(2.2), Inches(0.95))
        badge.fill.solid()
        badge.fill.fore_color.rgb = PRIMARY if i < 3 else SECONDARY
        badge.line.fill.background()

        tb_b = s7.shapes.add_textbox(Inches(0.85), top + Inches(0.25), Inches(2.1), Inches(0.5))
        tf_b = tb_b.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = fase
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = CARD_BG
        p_b.alignment = PP_ALIGN.CENTER

        tb_t = s7.shapes.add_textbox(Inches(3.2), top + Inches(0.08), Inches(9.1), Inches(0.8))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK

        p_d = tf_t.add_paragraph()
        p_d.text = desc.replace("\n", "  |  ")
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = MUTED

    # ==========================================
    # SLIDE 8: ESTRATÉGIA COM AS CLIENTES PILOTO
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Estratégia de Validação com as Clientes Piloto")

    piloto_cards = [
        ("1. Seleção (3 a 5 Confeiteiras)", 
         "• 2 que fazem bolos decorados/encomendas de festa.\n• 2 que fazem doces para pronta entrega (copos da felicidade, fatias).\n• 1 que já tem um ajudante ou atendente.\n\nEssa variedade valida tanto a parte de produção caseira quanto a equipe inicial.",
         PRIMARY),
        ("2. O Teste do 'Uau' no Dia 1", 
         "Cadastre junto com a cliente piloto a receita principal dela.\n\nQuando ela ver na tela: 'Seu bolo custa R$ 38,50 para produzir e você vendia a R$ 50,00, sobrando só R$ 11,50', ela ficará impressionada com a clareza do sistema.",
         ACCENT_GOLD),
        ("3. Métricas de Sucesso do Piloto", 
         "• Redução de 70% no tempo gasto explicando cardápio no WhatsApp.\n• 100% dos pedidos do fim de semana organizados sem esquecimento.\n• Segurança total de que nenhuma encomenda foi entregue sem receber.",
         SUCCESS),
    ]

    for i, (title, desc, col) in enumerate(piloto_cards):
        left = Inches(0.8 + i * 4.0)
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        top_bar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = col
        top_bar.line.fill.background()

        tb = s8.shapes.add_textbox(left + Inches(0.25), Inches(2.3), Inches(3.2), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = DARK
        p_d.space_before = Pt(12)

    # Save
    out_dir = r"Documentos"
    os.makedirs(out_dir, exist_ok=True)
    pptx_path = os.path.join(out_dir, "DoceGestor_Apresentacao_Piloto.pptx")
    prs.save(pptx_path)
    print(f"Apresentação PowerPoint salva com sucesso em: {pptx_path}")

create_presentation()
